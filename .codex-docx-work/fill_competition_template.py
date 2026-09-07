from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


REFERENCE = Path(r"C:\Users\HUAWEI\Documents\xwechat_files\wxid_wisnwe00gep312_e50a\msg\file\2026-09\附件2：2026年安徽省AI大模型创新应用竞赛作品赛竞赛模板.docx")
OUTPUT = Path(r"C:\Users\HUAWEI\Documents\ChatGPT\网站设计\中国色彩数字博物馆_作品报告书.docx")


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, size=10.5, bold=False, align=None, first_line=True):
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.space_before = Pt(0)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_paragraph_text(paragraph, text, size=10.5, bold=False, align=None, first_line=True):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    format_paragraph(paragraph, size=size, bold=bold, align=align, first_line=first_line)


def set_cell(cell, text, bold=False, center=False):
    cell.text = ""
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        run = paragraph.add_run(line)
        set_run_font(run, size=10.5, bold=bold)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def add_body_paragraph(doc, text="", heading=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, bold=heading)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(4 if not heading else 6)
    paragraph.paragraph_format.space_before = Pt(6 if heading else 0)
    paragraph.paragraph_format.first_line_indent = Pt(0 if heading else 21)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


doc = Document(REFERENCE)

for paragraph in doc.paragraphs:
    if paragraph.text.startswith("作品名称："):
        set_paragraph_text(paragraph, "作品名称：中国色彩数字博物馆", size=14, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)

table = doc.tables[0]
rows = table.rows

work_intro = (
    "中国色彩数字博物馆是一座面向公众的传统色彩知识与内容创作型数字展馆。作品以“色载华夏，数藏千年”为叙事线索，"
    "将青、赤、黄、白、黑五正色、器物色、节气色与可检索色彩档案组织为可浏览、可交互、可复用的网页体验。"
    "作品采用静态 HTML、CSS、JavaScript 实现前端展示，并以大模型辅助完成知识梳理、内容结构化、展陈文案生成、"
    "提示词设计和视觉资产策划，使分散的传统色知识转化为清晰的数字展览。用户可以通过序厅进入四个展厅，"
    "浏览五色体系，观察器物中的色彩，沿二十四节气理解时间中的颜色，并在拾色藏库中搜索、筛选、复制和组合传统色。"
)

set_cell(rows[0].cells[1], "中国色彩数字博物馆")
set_cell(rows[1].cells[1], "大模型赋能的知识管理与内容创作")
set_cell(rows[2].cells[0], "作品简介", bold=True, center=True)
try:
    rows[3].cells[0].merge(rows[3].cells[1])
except Exception:
    pass
set_cell(rows[3].cells[0], work_intro)
set_cell(rows[4].cells[1], "自主设计开发")
set_cell(rows[5].cells[1], "文化数字化、博物馆展陈、数字文旅、知识管理与内容创作")
set_cell(rows[6].cells[1], "公开传统色文化资料、五色体系与二十四节气知识、传统器物与工艺常识，以及作品设计期自建的色彩、展项和文案数据；不包含学校、团队、个人或用户隐私数据。")
set_cell(rows[7].cells[1], "本地静态网站程序：从项目根目录打开 index.html，或启动静态服务后访问序厅及四个展厅页面。")

for row in rows:
    shade_cell(row.cells[0], "F2F2F2")
    set_cell(row.cells[0], row.cells[0].text, bold=True, center=True)

set_cell(rows[2].cells[0], "作品简介", bold=True, center=True)
set_cell(rows[3].cells[0], work_intro, bold=False, center=False)

body_start = None
for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.strip() == "正文":
        body_start = index
        break

if body_start is not None:
    for paragraph in list(doc.paragraphs)[body_start + 1 :]:
        delete_paragraph(paragraph)

body_content = [
    ("选题定位", True),
    ("传统色是理解中华文化的重要入口，但相关知识长期分散在文献、器物、绘画、服饰、建筑、节令和设计语境之中。普通公众往往只能接触到单个色名或色卡，难以理解颜色背后的方位、五行、礼制、材质、工艺与生活场景；内容创作者也常面临资料零散、叙事结构难搭建、视觉表达难统一的问题。本作品选择“中国色彩数字博物馆”作为应用场景，尝试用大模型把传统色知识整理成可检索、可叙事、可交互的数字内容体系。", False),
    ("作品不是单纯展示颜色列表，而是把颜色作为信息架构和交互语言：用户先在序厅感知五色初醒，再进入“正色溯源”“物色生香”“岁时流转”“拾色藏库”四个展厅，从天地秩序、器物材料、四时节气和色彩档案四个层面理解中国传统色。", False),
    ("社会价值", True),
    ("作品面向传统文化学习者、设计爱好者、数字文旅体验用户、博物馆展陈策划者和内容创作者。对公众而言，它把抽象的传统色知识转化为可浏览的数字展览，降低理解门槛；对教育与文旅场景而言，它可作为传统文化课程、线上展陈和活动传播的轻量化原型；对设计与内容生产场景而言，它提供色名、色值、文化印象、来源语境和配色组合等素材，帮助创作者更准确地使用传统色。", False),
    ("在推广价值上，作品采用静态网页架构，部署成本低，便于在校园展示、文化活动、课程教学和线上传播中复用。其方法也可以扩展到非遗技艺、地域纹样、传统器物、地方文献等其他文化知识主题。", False),
    ("技术方案", True),
    ("作品采用前端静态网页实现，核心技术包括 HTML5 语义结构、CSS3 响应式布局与动效、原生 JavaScript 数据驱动交互。当前项目包含序厅入口和四个展厅页面：序厅负责建立整体叙事和导航；正色溯源展示青、赤、黄、白、黑与方位、五行、季节之间的关系；物色生香通过器物、服饰、书画和建筑呈现颜色依附材料与工艺的方式；岁时流转以二十四节气为时间轴组织颜色；拾色藏库提供搜索、筛选、排序、复制、收藏、经典配色和自定义配色组合。", False),
    ("大模型主要用于知识管理和内容创作两个环节。在知识管理环节，先将传统色相关资料拆分为色名、色系、色值、来源、应用、文化印象、描述文案、展厅归属等字段，形成可被前端读取的数据结构；再通过大模型辅助归纳颜色之间的关系，建立“五色体系-器物应用-节气场景-色彩档案”的内容层级。在内容创作环节，大模型用于生成展厅叙事、按钮与提示文案、色彩解释、器物导览语、页面标题和视觉资产提示词，并在人工审核后写入网页。", False),
    ("功能模块设计", True),
    ("序厅模块承担用户进入作品的第一层体验，通过视频背景、五色初醒、今日国色和四个展厅入口建立整体世界观。用户不需要先阅读复杂说明，即可通过滚动、点击五色按钮和随机换色理解作品的基本交互方式。正色溯源模块把五正色放在方位、五行和四时关系中展示，强调颜色背后的文化秩序，而不是只呈现视觉色块。物色生香模块把颜色放回器物、服饰、书画和建筑等载体，用户可以通过器物图像和色点交互理解颜色如何依附于材料、工艺与生活经验。", False),
    ("岁时流转模块采用二十四节气时间轴，将一年中的气候、物候与代表色进行策展式关联，使用户沿着时间顺序观察颜色变化。拾色藏库模块则承担知识管理结果的集中呈现，提供色谱总览、关键词搜索、色系筛选、朝代与应用筛选、明度和色相排序、颜色详情、色值复制、个人拾色收藏、经典配色和自定义配色等功能。四个展厅之间通过统一导航、颜色转场和页脚路径连接，形成从认知、观察、时间到归档的完整参观流程。", False),
    ("知识管理流程", True),
    ("作品将传统色知识处理为“采集、拆解、归类、生成、审核、发布”的流程。首先根据传统色名、五色体系、二十四节气、器物材料和应用场景收集公开资料；其次把自然语言资料拆解为结构化字段，例如色名、拼音、HEX 值、RGB 值、所属色系、文化来源、常见应用、象征意义和说明文案；然后由大模型辅助识别颜色之间的层级关系，把单个颜色归入五色体系、器物场景、节气场景或藏库档案。", False),
    ("在内容生成阶段，大模型并不直接决定最终结论，而是提供可编辑的初稿：例如为同一颜色生成通俗解释、展陈说明、按钮短语和视觉提示词。人工再根据准确性、语气、页面长度和交互位置进行筛选修改。最终数据以 JavaScript 数据文件的形式进入网页，前端根据字段自动生成色卡、筛选项、详情抽屉、配色组合和状态反馈。这样的流程使作品具备继续扩充的可能，后续新增颜色或展项时，只需要补充统一字段，页面即可按既有规则呈现。", False),
    ("用户体验设计", True),
    ("作品的交互设计遵循“少说明、多引导”的原则。页面中的导航、按钮、状态文字和过渡动效均围绕“观色、入馆、拾色、藏色”展开，尽量让用户从动作中理解内容。桌面端提供克制的自定义光标和滚动揭示，增强数字博物馆的沉浸感；移动端则保留清晰的导航、触控目标和垂直浏览节奏，确保在手机上也能完成浏览、筛选和复制。", False),
    ("在信息层级上，作品先给出感性的颜色和展厅名称，再逐步呈现更具体的知识字段。用户可以停留在短路径中完成一次轻量参观，也可以进入藏库进行更细的检索和配色组合。复制色值、收藏颜色和切换视图等功能把文化展示与实际创作需求连接起来，使用户不仅“看见颜色”，也能把颜色带入自己的设计、课程展示或内容创作。", False),
    ("创新性", True),
    ("与常见色卡网站相比，作品把大模型生成能力嵌入策展流程，而不是只把模型作为聊天入口。它用模型帮助完成资料整理、语义分类、展项命名、导览文案和交互提示，使文化知识从“资料堆叠”变成“展览叙事”。同时，作品把传统色与网页交互深度绑定，颜色不仅用于装饰，也参与导航、转场、筛选、状态反馈和内容组织，形成以色彩为核心的数字博物馆体验。", False),
    ("作品的另一项创新在于把传统文化内容生产拆成可复用的模型协作链路。传统文化类网页常见的问题是内容有文化气质但结构松散，或者数据可检索但缺少叙事温度。本作品以大模型辅助搭建“知识字段-展陈文案-视觉资产-前端组件”的连续流程，让知识管理成果能够直接转化为网页内容、交互状态和视觉表达。该流程既服务当前的中国传统色主题，也为其他文化知识数字化提供了可迁移的方法。", False),
    ("提示词设计及处理方法", True),
    ("提示词设计采用分层方法。第一层要求模型依据主题提取知识字段，输出结构化色彩档案；第二层要求模型按展厅定位生成策展叙事，控制语言为克制、准确、具有博物馆感；第三层要求模型面向网页组件生成短文案、交互提示和空状态提示，保证用户能在不阅读说明书的情况下完成浏览；第四层要求模型生成视觉资产和动效方向，再由人工筛选、修订和前端实现。", False),
    ("为避免内容失真，作品在处理时区分事实性知识、策展性表达和设计期创作内容。涉及历史制度、器物材料和传统色来源的文字以保守表述为主；涉及色值的数据明确定位为屏幕展示和设计参考，不把单一 HEX 值说成传统颜色的唯一标准。", False),
    ("具体处理时，提示词会先限定输出边界，例如要求模型不得编造馆藏编号、不得声称某一色值为历史唯一标准、不得出现参赛主体信息；再要求模型以表格或 JSON 形式输出字段，便于前端读取和人工校对。对于展陈文案，提示词会约束字数、语气、受众和使用位置，例如首页标题需要短而有记忆点，展厅说明需要有文化解释，按钮与空状态需要直接说明用户下一步动作。通过这种方式，大模型输出被约束在明确的内容生产任务中，降低随机性。", False),
    ("作品还采用“生成后校正”的方法处理大模型文本。对涉及传统制度、器物材料和节气知识的表述，优先选择谨慎、可验证、不过度绝对化的语言；对面向用户的互动文案，则强调清楚、简短和符合页面情境。这样既发挥模型在快速组织内容、提供表达方案方面的优势，也保留人工对事实边界和审美质量的把关。", False),
    ("实现与运行环境", True),
    ("作品前端文件由 index.html、halls 目录下的展厅页面、css 样式文件、js 交互脚本、data 数据文件和 assets 视觉资产组成。数据与页面分离，颜色、节气、器物和配色信息集中存放，页面脚本负责读取并渲染相应模块。该结构有利于维护，也方便后续将更多数据迁移到接口或数据库中。当前版本不依赖登录、后端服务或复杂构建环境，既可以通过本地文件浏览，也可以部署到普通静态站点托管平台。", False),
    ("在可访问性与稳定性方面，页面使用语义化标题、真实链接和按钮、可见焦点状态、复制反馈、无障碍标签和响应式布局。动效设计保持节制，并兼顾降低运动偏好的用户。由于作品定位为数字博物馆，视觉上避免过度娱乐化，主要通过留白、层级、图像、色块、转场和节奏建立东方审美与现代界面的结合。", False),
    ("应用效果", True),
    ("作品已经形成可运行的多页面网页原型。用户可以在首页通过沉浸式序厅进入展馆，在正色溯源页面理解五正色秩序，在物色生香页面通过器物图像与色点发现颜色，在岁时流转页面沿二十四节气浏览一年之色，在拾色藏库页面检索颜色、查看详情、复制色值、收藏颜色并组合配色。页面支持桌面端和移动端访问，并包含滚动揭示、页面转场、状态提示、复制反馈和无障碍语义标签等交互细节。", False),
    ("从展示效果看，作品完成了从单页介绍到多展厅体验的扩展。用户进入首页后能够获得明确的品牌识别和参观路径；进入正色溯源后，可以把五正色与传统秩序联系起来；进入物色生香后，可以看到颜色与材料、工艺、器物之间的关系；进入岁时流转后，可以沿时间轴建立四时观色的体验；进入拾色藏库后，可以把前面分散出现的颜色重新作为档案检索和组合。整个流程形成较完整的“认知-体验-检索-创作”闭环。", False),
    ("从内容创作效果看，大模型帮助作品在较短时间内完成了大量页面文案、交互文案和数据说明的统一。不同页面虽然承担不同功能，但语气保持为克制、准确、具有博物馆感，避免出现风格割裂。作品还将色值复制、个人拾色、经典配色和自定义组合等功能纳入藏库，使传统色知识不止停留在欣赏层面，也能成为后续设计和传播中的可用素材。", False),
    ("推广与应用场景", True),
    ("作品可应用于文化课程导入、博物馆线上导览、传统节气主题活动、数字文旅展示和设计教学。教师可以用它讲解五色体系、二十四节气与传统审美之间的关系；文旅或展陈策划者可以参考其展厅结构，把地方色彩、器物、建筑或非遗工艺整理成类似的数字专题；设计学习者则可以从藏库中寻找配色灵感，并理解颜色背后的文化语境。", False),
    ("若进一步完善数据来源和交互能力，作品还可以扩展为传统色知识管理平台。平台可支持多来源文献标注、图像色彩提取、颜色相似度检索、配色方案导出、面向儿童或专业设计师的不同讲解模式，以及大模型问答导览。届时用户不仅可以浏览预设展览，也可以围绕某个色名、节气、器物或应用场景生成个性化学习内容。", False),
    ("问题与改进方向", True),
    ("当前作品仍处于原型阶段，主要问题包括传统色资料来源需要进一步校核，色值与历史材料之间仍需更严谨的证据链，部分视觉资产属于设计期合成图，尚未接入权威馆藏图像。后续可进一步引入经过标注的文献与图像资料，建立可追溯的数据来源；增加大模型问答导览、自动生成学习卡片、面向不同年龄层的讲解模式，以及面向设计师的配色导出功能，使作品从展示型网页升级为可持续扩展的传统色知识管理平台。", False),
    ("后续迭代还需要加强评价机制。可邀请不同类型用户进行体验测试，观察他们是否能顺利完成入馆、筛选、复制和组合配色等任务；同时建立内容审核表，对色名、来源、应用、文化解释和生成文案逐项校对。通过用户反馈和内容校核双线推进，作品才能在保持审美体验的同时，提高文化知识表达的可靠性。", False),
    ("总结", True),
    ("总体而言，中国色彩数字博物馆以大模型辅助传统色知识管理和内容创作为核心，把分散的文化资料整理为可浏览、可检索、可交互的数字展陈。作品以网页原型验证了传统文化知识从结构化整理到策展表达、再到创作复用的完整路径，具有继续扩展为文化知识管理平台的基础。", False),
]

for text, heading in body_content:
    add_body_paragraph(doc, text, heading=heading)

for paragraph in doc.paragraphs:
    if paragraph.text.strip() in {"作品简介", "正文"}:
        format_paragraph(paragraph, size=12, bold=True, first_line=False)
    elif paragraph.text.strip() and paragraph.text not in {"填写说明"}:
        for run in paragraph.runs:
            if run.text:
                set_run_font(run, size=run.font.size.pt if run.font.size else 10.5, bold=run.bold)

doc.save(OUTPUT)
print(OUTPUT)
